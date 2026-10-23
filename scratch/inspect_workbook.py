import openpyxl
import sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'C:\Users\sable\Downloads\Sales_MIS_Validation_6000_Transactions_Executive_Sales_Dashboard_2026_CORRECTED.xlsx', data_only=False)
dash = wb['Executive Dashboard']
print('=== EXECUTIVE DASHBOARD (Formulas) ===')
for cell in ['B7', 'D7', 'F7', 'H7', 'J7', 'L7', 'N7']:
    print(f'{cell}: val={dash[cell].value}, num_format={dash[cell].number_format}')

wb_data = openpyxl.load_workbook(r'C:\Users\sable\Downloads\Sales_MIS_Validation_6000_Transactions_Executive_Sales_Dashboard_2026_CORRECTED.xlsx', data_only=True)
dash_data = wb_data['Executive Dashboard']
print('\n=== EXECUTIVE DASHBOARD (Cached Data) ===')
for cell in ['B7', 'D7', 'F7', 'H7', 'J7', 'L7', 'N7']:
    print(f'{cell}: result={dash_data[cell].value}')

piv = wb['Pivot Table Analysis']
print('\n=== PIVOT TABLE 1 (Headers) ===')
for c in ['B6', 'C6', 'D6', 'E6', 'F6', 'G6', 'H6', 'I6']:
    print(f'{c}: {piv[c].value}')

piv_data = wb_data['Pivot Table Analysis']
print('\n=== PIVOT TABLE 1 (Cached Row 7 - Region 1) ===')
for c in ['B7', 'C7', 'D7', 'E7', 'F7', 'G7', 'H7', 'I7']:
    print(f'{c}: formula={piv[c].value} -> cached={piv_data[c].value}')

print('\n=== PIVOT TABLE 1 (Total Row 11) ===')
for c in ['B11', 'G11', 'H11', 'I11']:
    print(f'{c}: label/formula={piv[c].value} -> cached={piv_data[c].value}')

tx = wb['Sales Transactions']
print('\n=== SALES TRANSACTIONS (Headers & Sample Row 3) ===')
for col in ['P', 'Q', 'R', 'S', 'T', 'U', 'V', 'W']:
    print(f'{col}2={tx[f"{col}2"].value} | {col}3 formula={tx[f"{col}3"].value} cached={wb_data["Sales Transactions"][f"{col}3"].value}')

reg = wb['Regional Performance']
print('\n=== REGIONAL PERFORMANCE (Headers & Row 3) ===')
print('E2:', reg['E2'].value, 'F2:', reg['F2'].value, 'G2:', reg['G2'].value, 'I2:', reg['I2'].value)
print('F3:', wb_data['Regional Performance']['F3'].value, 'I3:', wb_data['Regional Performance']['I3'].value, 'J3 formula:', reg['J3'].value, 'cached J3:', wb_data['Regional Performance']['J3'].value)

cust = wb['Customer Accounts']
print('\n=== CUSTOMER ACCOUNTS (Header Col 9 & Rows Count) ===')
print('I2:', cust['I2'].value)
print('Max Rows:', cust.max_row)
