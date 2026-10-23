import openpyxl, re, sys
sys.stdout.reconfigure(encoding='utf-8')

wb_path = sys.argv[1] if len(sys.argv) > 1 else r'C:\Users\sable\Downloads\Sales_MIS_Validation_6000_Transactions_Executive_Sales_Dashboard_2026 (1).xlsx'
print(f'Checking workbook: {wb_path}')
wb = openpyxl.load_workbook(wb_path, data_only=False)

deps = {}
cell_ref_regex = re.compile(r"(?:(?:'([^']+)'|([A-Za-z0-9_]+))!)?(\$?[A-Za-z]+\$?[0-9]+)")

for sheet in wb.sheetnames:
    ws = wb[sheet]
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                src = f'{sheet}!{cell.coordinate}'
                f = cell.value
                targets = []
                for m in cell_ref_regex.finditer(f):
                    s = m.group(1) or m.group(2) or sheet
                    c = m.group(3).replace('$', '')
                    targets.append(f'{s}!{c}')
                deps[src] = targets

print(f'Total formula cells: {len(deps)}')

state = {}
parent = {}
cycles = []
for node in deps:
    if state.get(node, 0) == 0:
        stack = [(node, iter(deps.get(node, [])))]
        state[node] = 1
        while stack:
            u, it = stack[-1]
            try:
                v = next(it)
                if v in deps:
                    v_state = state.get(v, 0)
                    if v_state == 1:
                        cycles.append((u, v))
                        print(f'CYCLE DETECTED: {u} -> {v}')
                    elif v_state == 0:
                        state[v] = 1
                        parent[v] = u
                        stack.append((v, iter(deps.get(v, []))))
            except StopIteration:
                state[u] = 2
                stack.pop()

print(f'Total cycles detected: {len(cycles)}')
