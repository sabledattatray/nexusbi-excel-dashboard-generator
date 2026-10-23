import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb_path = r'C:\Users\sable\Downloads\Sales_MIS_Validation_6000_Transactions_Executive_Sales_Dashboard_2026 (1).xlsx'
wb = openpyxl.load_workbook(wb_path, data_only=False)
wb_data = openpyxl.load_workbook(wb_path, data_only=True)

print(f"Sheet names ({len(wb.sheetnames)}): {wb.sheetnames}")

# 1. Executive Dashboard
ws_dash = wb['Executive Dashboard']
ws_dash_val = wb_data['Executive Dashboard']

print("\n--- Executive Dashboard Cards ---")
print("B6 (Label):", ws_dash['B6'].value, "| B7 (Formula):", ws_dash['B7'].value, "| B7 (Val):", ws_dash_val['B7'].value)
print("D6 (Label):", ws_dash['D6'].value, "| D7 (Formula):", ws_dash['D7'].value, "| D7 (Val):", ws_dash_val['D7'].value)
print("F6 (Label):", ws_dash['F6'].value, "| F7 (Formula):", ws_dash['F7'].value, "| F7 (Val):", ws_dash_val['F7'].value)
print("H6 (Label):", ws_dash['H6'].value, "| H7 (Formula):", ws_dash['H7'].value, "| H7 (Val):", ws_dash_val['H7'].value)
print("J6 (Label):", ws_dash['J6'].value, "| J7 (Formula):", ws_dash['J7'].value, "| J7 (Val):", ws_dash_val['J7'].value)
print("L6 (Label):", ws_dash['L6'].value, "| L7 (Formula):", ws_dash['L7'].value, "| L7 (Val):", ws_dash_val['L7'].value)
print("N6 (Label):", ws_dash['N6'].value, "| N7 (Formula):", ws_dash['N7'].value, "| N7 (Val):", ws_dash_val['N7'].value)

# 2. Sales Transactions
ws_tx = wb['Sales Transactions']
ws_tx_val = wb_data['Sales Transactions']
max_r = ws_tx.max_row
print(f"\n--- Sales Transactions (Rows: {max_r}) ---")
headers = [ws_tx.cell(2, c).value for c in range(1, 30)]
print("Headers (Row 2, Cols 1-28):", headers[:28])
print("Row 3 sample:", [ws_tx_val.cell(3, c).value for c in [1, 2, 4, 8, 12, 14, 18, 19, 20, 21, 22, 23, 26, 27, 28]])
print("Total Row (Row 6003):")
print("  Qty (Col R):", ws_tx['R6003'].value, "| Val:", ws_tx_val['R6003'].value)
print("  Net Sales (Col T):", ws_tx['T6003'].value, "| Val:", ws_tx_val['T6003'].value)
print("  Total COGS (Col U):", ws_tx['U6003'].value, "| Val:", ws_tx_val['U6003'].value)
print("  Gross Profit (Col V):", ws_tx['V6003'].value, "| Val:", ws_tx_val['V6003'].value)
print("  Margin % (Col W):", ws_tx['W6003'].value, "| Val:", ws_tx_val['W6003'].value)
print("  Gross Sales (Col Z):", ws_tx['Z6003'].value, "| Val:", ws_tx_val['Z6003'].value)
print("  Discount Amt (Col AA):", ws_tx['AA6003'].value, "| Val:", ws_tx_val['AA6003'].value)

# 3. Regional Performance
ws_reg = wb['Regional Performance']
ws_reg_val = wb_data['Regional Performance']
print(f"\n--- Regional Performance (Rows: {ws_reg.max_row}) ---")
print("Headers:", [ws_reg.cell(2, c).value for c in range(1, 11)])
for r in range(3, ws_reg.max_row + 1):
    vals = [ws_reg_val.cell(r, c).value for c in range(1, 11)]
    formulas = [ws_reg.cell(r, c).value if str(ws_reg.cell(r, c).value).startswith('=') else '' for c in range(1, 11)]
    print(f"  Row {r}: {vals[0]} | Orders: {vals[2]} | Units: {vals[3]} | Rev: {vals[4]} | Cost: {vals[5]} | Profit: {vals[6]} | Margin: {vals[7]} | Share: {vals[8]}")

# 4. Product Catalog
ws_prod = wb['Product Catalog']
ws_prod_val = wb_data['Product Catalog']
print(f"\n--- Product Catalog (Rows: {ws_prod.max_row}) ---")
print("Headers:", [ws_prod.cell(2, c).value for c in range(1, 13)])
print("Sample Row 3:", [ws_prod_val.cell(3, c).value for c in range(1, 13)])
print("Total Row:", [ws_prod_val.cell(ws_prod.max_row, c).value for c in range(1, 13)])

# 5. Customer Accounts
ws_cust = wb['Customer Accounts']
ws_cust_val = wb_data['Customer Accounts']
print(f"\n--- Customer Accounts (Rows: {ws_cust.max_row}) ---")
print("Headers:", [ws_cust.cell(2, c).value for c in range(1, 13)])
print("Sample Top Customer (Row 3):", [ws_cust_val.cell(3, c).value for c in range(1, 13)])
print("Remaining Accounts Row:", [ws_cust_val.cell(ws_cust.max_row - 1, c).value for c in range(1, 13)])
print("Grand Total Portfolio Row:", [ws_cust_val.cell(ws_cust.max_row, c).value for c in range(1, 13)])

# 6. Audit & Data Provenance
ws_aud = wb['Audit & Data Provenance']
ws_aud_val = wb_data['Audit & Data Provenance']
print(f"\n--- Audit & Data Provenance (Rows: {ws_aud.max_row}) ---")
for r in range(4, 15):
    print(f"  Sec 1 Row {r}:", ws_aud_val.cell(r, 1).value, "->", ws_aud_val.cell(r, 4).value, "(", ws_aud_val.cell(r, 8).value, ")")

print("\nCanonical Metric Registry sample:")
for r in range(18, 31):
    print(f"  Sec 2 Row {r}:", ws_aud_val.cell(r, 1).value, "|", ws_aud_val.cell(r, 2).value, "|", ws_aud_val.cell(r, 6).value, "|", ws_aud_val.cell(r, 7).value, "|", ws_aud_val.cell(r, 8).value)

print("\nCross-Dimensional Reconciliation Matrix:")
for r in range(33, 42):
    print(f"  Sec 3 Row {r}:", ws_aud_val.cell(r, 1).value, "|", ws_aud_val.cell(r, 3).value, "| Rev:", ws_aud_val.cell(r, 4).value, "| Cost:", ws_aud_val.cell(r, 5).value, "| Profit:", ws_aud_val.cell(r, 6).value, "| Var:", ws_aud_val.cell(r, 8).value)
